from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
OUTPUT_PATH = DOCS_DIR / "restructured_lms_workflow_v4.docx"


@dataclass
class TextSection:
    number: str
    title: str
    paragraphs: list[str]
    bullets: list[str]


@dataclass
class ScreenSection:
    number: str
    title: str
    image: Path
    what_happens: list[str]
    on_submit: list[str]
    approval_flow: list[str]
    api_communication: list[str]


def screen(number: str, title: str, image: str, what: list[str], on_submit: list[str],
           approval: list[str], api: list[str]) -> ScreenSection:
    return ScreenSection(
        number=number,
        title=title,
        image=(ROOT / image).resolve(),
        what_happens=what,
        on_submit=on_submit,
        approval_flow=approval,
        api_communication=api,
    )


STATUS_REPLACEMENTS = {
    "`ACTIVE`": "active",
    "`NEW`": "new",
    "`DRAFT`": "Draft",
    "`APPROVED`": "approved",
    "`PENDING`": "pending",
    "`AWAITING_GUARANTORS`": "Awaiting Guarantors",
    "`ALL_GUARANTORS_APPROVED`": "All Guarantors Approved",
    "`READY_FOR_MANAGER`": "Ready For Manager",
    "`ON REVIEW BY MANAGER`": "On Review By Manager",
    "`AWAITING_BOARD`": "On Review By Board",
    "`BOARD_APPROVED`": "Board Approved",
    "`BOARD_REJECTED`": "Board Rejected",
    "`MANAGER_REJECTED`": "Manager Rejected",
    "`FINAL_APPROVED`": "Final Approved",
    "`FINAL_REJECTED`": "Final Rejected",
    "`DISBURSED LOAN`": "Disbursed Loan",
    "`PAID`": "Paid",
}


USER_COPY_OVERRIDES = {
    "Spring Security receives staff password login through `/login`, while the OTP path is reserved for staff accounts that already exist locally and are marked `ACTIVE`.":
        "The page lets staff sign in with a password or use a one-time code when that option is available for their account.",
    "Password login delegates to `SecurityConfig` and `AppUserDetailsService`, then redirects by the authenticated staff member's primary role.":
        "If the password is correct, the user is signed in and taken to the correct work area.",
    "OTP login first requests a code, then verifies it, then writes the authenticated `AppUserPrincipal` into the HTTP session.":
        "For one-time-code sign-in, the system sends the code, checks it, and then signs the user in.",
    "The next visible step is workspace routing to `/admin/dashboard`, `/manager/loan-applications?status=READY_FOR_MANAGER`, or `/board/queue` depending on staff role.":
        "After sign-in, the user is taken to the correct dashboard or review queue based on their role.",
    "The screen is the only place where a member chooses `saccoId` and `stationId`; after registration, those values are fixed on the `Member` row and reused from the authenticated principal.":
        "This is where the member chooses the correct SACCO and station, and those details stay linked to the account after registration.",
    "The successful outcome is an `ACTIVE` member account that can later sign in and start loan applications within the registered SACCO only.":
        "The result is an active member account that can sign in and apply for loans in the selected SACCO.",
    "The OTP request step validates the form bean, verifies the external member profile, checks that the station belongs to the selected SACCO, and blocks duplicate local registration.":
        "When the member requests a one-time code, the system checks the entered details, confirms the member belongs to the selected SACCO and station, and stops duplicate registration.",
    "The final submit step consumes the registration OTP, creates the local member account, creates savings and user-settings rows, and redirects the user back to login.":
        "The final submit step confirms the one-time code, creates the member account, prepares the supporting account details, and returns the user to the login page.",
    "For platform administrators, `/admin/dashboard` renders the platform dashboard rather than the workspace dashboard.":
        "This dashboard gives platform administrators a full view across all SACCO workspaces.",
    "The cards summarize registered SACCOs, active members, total disbursed principal, failed outbox events, and recent platform activity across every active SACCO.":
        "The cards summarize registered SACCOs, active members, total disbursed amounts, failed message deliveries, and recent platform activity across every active SACCO.",
    "It is the operational launch point into SACCO portfolio management, registry maintenance, outbox inspection, and audit history.":
        "It is the starting point for SACCO portfolio management, registry maintenance, delivery monitoring, and activity history.",
    "Clicking `Add SACCO` opens the creation modal, and row actions open detail or edit flows.":
        "Selecting Add SACCO opens the form for creating a new workspace, while the row actions open details or editing options.",
    "Station rows are activated and local SACCO settings are created or updated at the same time.":
        "The listed stations are activated and the SACCO's basic settings are prepared at the same time.",
    "The form binds each invited staff member to one SACCO and one station, which keeps workspace administration scoped cleanly.":
        "The form links each invited staff member to one SACCO and one station so their access stays limited to the correct workspace.",
    "The outbox monitor lists platform events that were queued by workflow actions such as guarantor assignment, board assignment, approvals, and disbursement updates.":
        "The outbox monitor lists system messages created by workflow actions such as guarantor assignment, board assignment, approvals, and disbursement updates.",
    "The actionable write on this screen is retrying a failed or blocked outbox event.":
        "The main action on this screen is retrying a failed or blocked message.",
    "The outbox does not create new loan decisions.":
        "This screen does not create new loan decisions; it only tracks message delivery.",
    "Retrying moves the selected event back to `NEW` so the publisher can process it again.":
        "Retrying returns the failed item to the queue so the system can try sending it again.",
    "The workspace dashboard shows the operational summary for the currently scoped SACCO and station.":
        "The workspace dashboard shows the operational summary for the currently selected SACCO and station.",
    "It routes the admin into incidents, users, settings, outbox, and event history for the selected workspace.":
        "It gives the admin quick access to incidents, users, settings, outbox, and event history for the selected workspace.",
    "Creating a user validates role assignment rules and persists a new local staff account in the current SACCO.":
        "Creating a user checks the selected role and then adds a new staff account to the current SACCO.",
    "The query and pagination controls help the admin manage a growing workspace user list without leaving the scoped SACCO.":
        "The search and page controls help the admin manage a growing user list within the current SACCO.",
    "The station registry shows the station set that belongs to the currently scoped SACCO.":
        "The station registry shows the stations that belong to the current SACCO.",
    "From here the admin can change the rules that `LoanWorkflowService.saveDraft(...)` and the eligibility preview later enforce.":
        "From here the admin can change the product rules that members must meet when preparing loan applications.",
    "The workspace outbox monitor is the scoped version of the platform outbox screen.":
        "The workspace outbox monitor is the SACCO-level version of the platform outbox screen.",
    "The outbox mirrors workflow transitions; it does not create them.":
        "This screen reflects workflow changes after they happen; it does not create them.",
    "It routes the member into product selection, application detail, archives, guarantees, reports, and support.":
        "It gives the member quick access to product selection, application detail, archives, guarantees, reports, and support.",
    "This quorum value is a business-critical setting because it determines when `AWAITING_BOARD` can transition to `BOARD_APPROVED` or `BOARD_REJECTED`.":
        "This setting is important because it decides how many board decisions are needed before a board result becomes final.",
    "It changes the future evaluation threshold that `BoardService.evaluateOutcome(...)` applies.":
        "Any change here affects future board reviews by changing how many decisions are needed.",
    "The product set is scoped by `principal.getSaccoId()` and not chosen freely by the member.":
        "Members only see the loan products available in their SACCO.",
    "Selecting Apply sends the member to the new application form for the chosen `loanType`.":
        "Selecting Apply opens the application form for the chosen product.",
    "The next workflow step is the product-aware draft form, where the first `DRAFT` row can be written.":
        "The next step is opening a new application draft for that product.",
    "The eligibility and example values are based on `EligibilityService` and `FinancialDetailsService`.":
        "The calculator shows an estimate based on the member's current savings position and the amount entered.",
    "When the full form is in use, the same preview pipeline is exposed through `GET /app/loan-applications/external-eligibility-summary` and `POST /app/loan-applications/financial-preview`.":
        "When the member opens the full application form, the same live preview is used to show current limits and repayment guidance.",
    "Saving writes or updates a `DRAFT` loan through `LoanWorkflowService.saveDraft(...)` after validating repayment months, dynamic schema fields, guarantor count, same-SACCO guarantors, and eligibility policy.":
        "Saving keeps the application as a draft after checking the repayment period, required form details, guarantor count, and current eligibility rules.",
    "The visible `Choose exactly N guarantor(s)` rule comes from the active product configuration for the member's SACCO.":
        "The required number of guarantors comes from the loan product rules set for that SACCO.",
    "This detail screen shows the application after all required guarantors have approved, so the status is `ALL_GUARANTORS_APPROVED` and the final member submission block is now unlocked.":
        "This screen shows the application after all required guarantors have approved, so the final submission to the manager is now unlocked.",
    "Submitting from this state validates the applicant OTP, records the applicant signature text and verification timestamp, and calls `LoanWorkflowService.submitToManager(...)`.":
        "Submitting here asks the member to confirm with a one-time code and then forwards the application to the manager.",
    "The visible next status is `READY_FOR_MANAGER`, which the UI labels as `ON REVIEW BY MANAGER`.":
        "After this step, the application moves to manager review.",
    "This detail state shows a member-owned application that is still `DRAFT`, with the declaration block and applicant OTP area visible before the application leaves the member workspace.":
        "This screen shows a member's application while it is still being prepared, before it leaves the member area.",
    "If the application remains a draft, `LoanWorkflowService.saveDraft(...)` can continue updating it; if the member submits, `LoanWorkflowService.submit(...)` decides whether the next state is guarantor review or manager review.":
        "While it is still a draft, the member can keep updating it. Once submitted, it moves either to guarantor review or straight to manager review, depending on the product rules.",
    "Once the loan moves to `AWAITING_GUARANTORS` or `READY_FOR_MANAGER`, the workflow becomes progressively more controlled and visible to other roles.":
        "Once the application leaves draft, it becomes more controlled and other reviewers can start seeing it.",
    "If the member clicks cancel while the loan is still at manager entry stage, the system can route the removal request through the reversal flow instead of silently discarding a review-stage item.":
        "If the member clicks cancel while the application is still at an early review step, the system handles that cancellation formally instead of removing it without trace.",
    "Outbox events are queued so the manager-side workflow can react to the handoff.":
        "A handoff message is prepared so the manager side can react to the submission.",
    "The member supplies subject and message text only; the SACCO context is inherited from the authenticated principal.":
        "The member only enters the subject and message, and the system links the request to the correct SACCO and member account automatically.",
    "Approval requires the guarantor to confirm the declaration, request an OTP, validate it, and then call `LoanWorkflowService.approveGuarantorRequest(...)` with signature text and verified time.":
        "To approve, the guarantor confirms the declaration, requests a one-time code, verifies it, and submits the decision.",
    "When all required guarantor requests reach `APPROVED`, the application moves from `AWAITING_GUARANTORS` to `ALL_GUARANTORS_APPROVED`.":
        "When all required guarantors approve, the application becomes ready for the member to send to the manager.",
    "The manager queue filtered to `ON REVIEW BY MANAGER` lists applications that have reached `READY_FOR_MANAGER` and are waiting for an accept or reject decision.":
        "This queue shows applications that have reached manager review and are waiting for a decision.",
    "The next possible states are `MANAGER_REJECTED` or `AWAITING_BOARD` after the accept path creates board assignments.":
        "From here, the manager can either reject the application or send it forward for board review.",
    "No decision is persisted until the manager posts accept or reject from the detail screen.":
        "No decision becomes final until the manager accepts or rejects the application from the detail screen.",
    "The next visible progression is from `AWAITING_BOARD` to either `BOARD_APPROVED` or `BOARD_REJECTED`, depending on quorum outcome.":
        "The next step is a board result, which becomes either approved or rejected once enough board decisions have been recorded.",
    "Rejecting writes a `ManagerReview` row, sets the loan to `MANAGER_REJECTED`, and queues a manager-rejected outbox event for the applicant.":
        "Rejecting records the manager's reason and marks the application as rejected.",
    "Accepting also writes a manager review row, temporarily sets `MANAGER_ACCEPTED`, assigns the required number of board reviewers, queues board-assignment outbox events, and finally moves the loan to `AWAITING_BOARD`.":
        "Accepting records the manager's decision, assigns the application to the board, and moves it into board review.",
    "The visible next stage after acceptance is `ON REVIEW BY BOARD`.":
        "After acceptance, the application appears as being on review by the board.",
    "Submitting final approval validates the disbursement date, first repayment date, unique numeric loan ID, and schedule inputs, then builds the repayment schedule and marks the loan `FINAL_APPROVED`.":
        "Final approval confirms the disbursement details, creates the repayment schedule, and marks the loan as approved for release.",
    "The next visible business state after the successful finalize action is the UI label `DISBURSED LOAN`.":
        "After successful final approval, the loan appears as a disbursed loan.",
    "The board queue lists the applications assigned to the signed-in board member that still have a `PENDING` board decision.":
        "This queue shows the applications assigned to the current board member that still need a decision.",
    "The board-stage application remains in `AWAITING_BOARD` until the configured quorum is met for approvals or rejections.":
        "The application stays in board review until enough board decisions have been recorded.",
    "When the review is still pending for the current assessor, the same route can expose the approve or reject form and the board OTP controls; once the decision is already recorded, the page becomes an audit-style detail view.":
        "When the current board member has not yet decided, this page shows the approve or reject form and the one-time-code steps. After a decision is recorded, the page becomes a read-only history view.",
    "The application remains in board workflow until `BoardService.evaluateOutcome(...)` sees enough approvals or rejections to satisfy quorum.":
        "The application stays in board review until enough approvals or rejections have been recorded to reach the required threshold.",
    "Once quorum is reached, the loan moves to `BOARD_APPROVED` or `BOARD_REJECTED`, and the manager later handles final approval or final rejection.":
        "Once that threshold is reached, the board result becomes final and the manager handles the last step.",
}


EXTERNAL_COMMUNICATION_BY_SECTION = {
    "2.2": [
        "During registration, the system checks the member's details against the Foresight financial app endpoint before the account is created.",
        "If the returned details do not match the form, registration is stopped and the member is asked to correct the information before continuing.",
    ],
    "5.4": [
        "While the member is filling this form, the system can pull current savings and related financial status from the Foresight financial app endpoint.",
        "That live check helps the form show the latest borrowing limit and prevents the member from continuing with outdated financial information.",
    ],
    "6.4": [
        "When the manager checks a guarantor's financial status, the latest summary is pulled from the Foresight financial app endpoint.",
        "This helps the manager review the guarantor's current position before making a decision.",
    ],
    "6.6": [
        "When payment information is refreshed, the system pulls the latest loan payments and repayment summary from the Members Portal endpoint.",
        "That returned information helps the system show the current repayment progress and decide whether the loan should remain active, move to Paid, or move to Defaulted.",
    ],
    "7.2": [
        "When a board member checks a guarantor's financial status, the latest summary is pulled from the Foresight financial app endpoint.",
        "This gives the board member up-to-date financial context before approving or rejecting the application.",
    ],
}


def simplify_user_copy(text: str) -> str:
    if text in USER_COPY_OVERRIDES:
        return USER_COPY_OVERRIDES[text]
    for old, new in STATUS_REPLACEMENTS.items():
        text = text.replace(old, new)
    return text.replace("`", "")


def build_sections() -> tuple[list[TextSection], list[tuple[str, str]], list[ScreenSection]]:
    intro_sections = [
        TextSection(
            number="1.1",
            title="Purpose And Scope",
            paragraphs=[
                "This document turns the raw LMS workflow reference into a clean, export-ready guide that explains what each screen is showing in simple user language.",
                "Each screenshot is matched with a clear workflow explanation so readers can understand what the user is doing, what the system is checking, and what happens next.",
            ],
            bullets=[
                "The document covers the shared login and registration pages, admin workspaces, member pages, manager pages, board pages, and printed output.",
                "Workflow names are kept consistent with the labels shown in the system so the explanations remain easy to follow.",
                "External-system notes are included only where the workflow depends on the Foresight financial app or the Members Portal.",
            ],
        ),
        TextSection(
            number="1.2",
            title="Workflow Reading Convention",
            paragraphs=[
                "Every screenshot is followed by a Workflow Explanation block with the same reading pattern: what the screen does, what happens when the user submits something, and how the process moves to the next step.",
            ],
            bullets=[
                "Read-only dashboards and lists are explained as viewing points in the workflow, even when no data is being changed.",
                "Action screens explain the visible checks first, then the result of the action, then the next review step where relevant.",
                "When a screenshot already shows a completed state, the explanation describes that completed state instead of inventing an extra action.",
            ],
        ),
    ]

    major_sections = [
        ("2", "Shared Entry Screens"),
        ("3", "Platform Administration (Super Admin)"),
        ("4", "Workspace Administration (Minor Admin)"),
        ("5", "Member Workflow Screens"),
        ("6", "Manager Workflow Screens"),
        ("7", "Board Workflow Screens"),
        ("8", "Printable Output"),
    ]

    screens = [
        screen(
            "2.1",
            "Staff Login",
            "docs/assets/lms-workflow-screens/shared/login-staff.png",
            [
                "The login page exposes the staff tab for password login and email OTP login on the same shared entry screen.",
                "Spring Security receives staff password login through `/login`, while the OTP path is reserved for staff accounts that already exist locally and are marked `ACTIVE`.",
            ],
            [
                "Password login delegates to `SecurityConfig` and `AppUserDetailsService`, then redirects by the authenticated staff member's primary role.",
                "OTP login first requests a code, then verifies it, then writes the authenticated `AppUserPrincipal` into the HTTP session.",
            ],
            [
                "No loan status changes happen on this screen.",
                "The next visible step is workspace routing to `/admin/dashboard`, `/manager/loan-applications?status=READY_FOR_MANAGER`, or `/board/queue` depending on staff role.",
            ],
            [
                "`POST /login` handles password login.",
                "`POST /login/staff/request-otp` issues the six-digit code through `EmailOtpService`.",
                "`POST /login/staff/verify-otp` validates the code, signs the user in, and returns the redirect URL for the correct staff workspace.",
            ],
        ),
        screen(
            "2.2",
            "Member Registration",
            "docs/assets/lms-architecure/screenshots/registration-page.png",
            [
                "The registration form captures member number, name, email, SACCO, station, and OTP so the member can be bound to one SACCO and one station at creation time.",
                "The screen is the only place where a member chooses `saccoId` and `stationId`; after registration, those values are fixed on the `Member` row and reused from the authenticated principal.",
            ],
            [
                "The OTP request step validates the form bean, verifies the external member profile, checks that the station belongs to the selected SACCO, and blocks duplicate local registration.",
                "The final submit step consumes the registration OTP, creates the local member account, creates savings and user-settings rows, and redirects the user back to login.",
            ],
            [
                "This is an onboarding step rather than a loan approval step.",
                "The successful outcome is an `ACTIVE` member account that can later sign in and start loan applications within the registered SACCO only.",
            ],
            [
                "`POST /register/member/request-otp` calls `MemberRegistrationService.verifyExternalMember(...)`, `ensureLocalUniqueness(...)`, and `EmailOtpService.issueOtp(...)`.",
                "`POST /register/member` consumes the OTP and calls `MemberRegistrationService.register(...)`.",
                "Registration validations include SACCO existence, valid station, external profile match, local uniqueness, and OTP validity.",
            ],
        ),
        screen(
            "3.1",
            "Super Admin Dashboard",
            "docs/assets/lms-workflow-screens/super-admin/dashboard.png",
            [
                "For platform administrators, `/admin/dashboard` renders the platform dashboard rather than the workspace dashboard.",
                "The cards summarize registered SACCOs, active members, total disbursed principal, failed outbox events, and recent platform activity across every active SACCO.",
            ],
            [
                "This dashboard is read-only and does not submit business data.",
                "It is the operational launch point into SACCO portfolio management, registry maintenance, outbox inspection, and audit history.",
            ],
            [
                "No loan application status is updated here.",
                "The screen supports oversight before a super admin drills into a specific SACCO or platform administration task.",
            ],
            [
                "`GET /admin/dashboard` routes Admin users to `PlatformAdminService.dashboard()`.",
                "The summary is assembled from `RegisteredSacco`, `Member`, `LoanApplication`, `SavingsAccount`, `AuditLog`, and `OutboxEvent` repositories.",
            ],
        ),
        screen(
            "3.2",
            "SACCOs Portfolio",
            "docs/assets/lms-workflow-screens/super-admin/saccos.png",
            [
                "The SACCO portfolio screen lists the active SACCO workspaces and highlights each portfolio's stations, members, exposure, and health indicators.",
                "This is the visual summary view of platform tenancy rather than the formal registry maintenance table.",
            ],
            [
                "The screen itself is read-only.",
                "The user moves from here into SACCO detail by opening a single workspace record.",
            ],
            [
                "No workflow state changes happen on the portfolio cards.",
                "The cards help platform staff decide which SACCO needs closer review before opening its detail page.",
            ],
            [
                "`GET /admin/saccos` returns `admin/platform-saccos` for Admin users.",
                "Each card is backed by `PlatformAdminService.dashboard()` portfolio summaries, not by member-side loan actions.",
            ],
        ),
        screen(
            "3.3",
            "SACCO Registry",
            "docs/assets/lms-workflow-screens/super-admin/sacco-registry.png",
            [
                "The SACCO registry is the formal list of registered SACCO workspaces and their station sets.",
                "Unlike the portfolio view, this table is a maintenance screen used to inspect or launch edit actions against registered SACCO metadata.",
            ],
            [
                "Clicking `Add SACCO` opens the creation modal, and row actions open detail or edit flows.",
                "No loan records are changed directly on the list page.",
            ],
            [
                "The approval chain is administrative rather than credit-related.",
                "The outcome of this page is a controlled set of active SACCO IDs and stations that later constrain registration, scoping, and reporting.",
            ],
            [
                "`GET /admin/saccos/registry` renders the registry for platform admins.",
                "The records come from `SaccoRegistryService.listRegisteredSaccos()` and active station rows.",
            ],
        ),
        screen(
            "3.4",
            "Add SACCO Modal",
            "docs/assets/lms-workflow-screens/super-admin/sacco-add-modal.png",
            [
                "The modal captures SACCO ID, SACCO name, station IDs, and an optional logo in one platform-admin action.",
                "It is the controlled entry point for creating a new workspace that the rest of the platform can reference safely.",
            ],
            [
                "Submitting the modal normalizes the SACCO ID and stations, requires at least one station, stores the logo if supplied, and creates default SACCO configuration.",
                "Station rows are activated and local SACCO settings are created or updated at the same time.",
            ],
            [
                "No loan status moves here.",
                "The important downstream effect is that member registration, admin scoping, and SACCO-specific loan configuration can now target the newly registered workspace.",
            ],
            [
                "`POST /admin/saccos` calls `SaccoRegistryService.registerSacco(...)`.",
                "The service persists `registered_saccos`, station rows, logo storage, and default `SaccoSettings`, then ensures default loan products exist for the new SACCO.",
            ],
        ),
        screen(
            "3.5",
            "SACCO Detail",
            "docs/assets/lms-workflow-screens/super-admin/sacco-detail.png",
            [
                "The SACCO detail page is the platform-level drill-down for one registered workspace.",
                "It summarizes member counts, disbursed exposure, paid and overdue loans, and recent audit activity for the selected SACCO.",
            ],
            [
                "This screenshot is a read-only overview state.",
                "Administrative edits are launched from the SACCO registry actions or update flows rather than typed directly into the overview cards.",
            ],
            [
                "No live approval action is taken here.",
                "The page helps platform administrators judge portfolio health before changing registry data or investigating workspace operations.",
            ],
            [
                "`GET /admin/saccos/{saccoId}` resolves `PlatformAdminService.saccoDetail(saccoId)`.",
                "The service joins registered SACCO data with member, loan, savings, and audit repositories to build the summary model.",
            ],
        ),
        screen(
            "3.6",
            "Minor Admin Registration",
            "docs/assets/lms-workflow-screens/super-admin/minor-admin-registration.png",
            [
                "This screen is used by the platform admin to invite or maintain Minor Admin accounts for a specific SACCO and station.",
                "The form binds each invited staff member to one SACCO and one station, which keeps workspace administration scoped cleanly.",
            ],
            [
                "Registering a Minor Admin validates the SACCO, validates the station, reserves a unique member number and email, and issues an activation link.",
                "Existing rows can also be updated, re-invited, revoked, or deactivated through the actions shown in the table.",
            ],
            [
                "The workflow outcome is administrative access, not a loan approval transition.",
                "Once active, the invited account can sign in and administer only its assigned workspace.",
            ],
            [
                "`GET /admin/saccos/minor-admins` renders the registration and management screen.",
                "`POST /admin/saccos/minor-admins` creates the invited account through `AdminService.registerMinorAdmin(...)`.",
                "Related actions use `/reinvite`, `/resend-invite`, `/revoke-invite`, and `/deactivate` routes for account lifecycle control.",
            ],
        ),
        screen(
            "3.7",
            "Platform Outbox Monitor",
            "docs/assets/lms-workflow-screens/super-admin/outbox.png",
            [
                "The outbox monitor lists platform events that were queued by workflow actions such as guarantor assignment, board assignment, approvals, and disbursement updates.",
                "The filter row lets administrators narrow by date range, page, and loan reference when tracing message delivery.",
            ],
            [
                "The actionable write on this screen is retrying a failed or blocked outbox event.",
                "Retrying moves the selected event back to `NEW` so the publisher can process it again.",
            ],
            [
                "The outbox does not create new loan decisions.",
                "It reflects already-recorded business events and helps operations verify that downstream notifications or integrations were queued correctly.",
            ],
            [
                "`GET /admin/outbox` returns filtered `OutboxEvent` pages.",
                "`POST /admin/outbox/{id}/retry` calls `AdminService.retryOutbox(...)` and resets the event for re-publication.",
            ],
        ),
        screen(
            "3.8",
            "Platform Event Log",
            "docs/assets/lms-workflow-screens/super-admin/events.png",
            [
                "The event log is the platform-wide audit trail of who did what, against which entity, and when it happened.",
                "It is used for traceability across member, manager, board, and administration actions.",
            ],
            [
                "This screen is read-only.",
                "Filter controls narrow the audit history by date and actor so support staff can reconstruct a workflow path.",
            ],
            [
                "No approval state changes occur here.",
                "The event log is evidence of transitions that were already committed elsewhere in the system.",
            ],
            [
                "`GET /admin/events` loads filtered audit pages through `AdminService.eventEntries(...)`.",
                "Entries originate from persisted `AuditLog` rows and include actions recorded by audit-aware services and controllers.",
            ],
        ),
        screen(
            "4.1",
            "Admin Dashboard",
            "docs/assets/lms-workflow-screens/minor-admin/dashboard.png",
            [
                "The workspace dashboard shows the operational summary for the currently scoped SACCO and station.",
                "Cards, database utilization, recent events, and recent incidents help the Minor Admin assess workspace health quickly.",
            ],
            [
                "The dashboard itself does not write business data.",
                "It routes the admin into incidents, users, settings, outbox, and event history for the selected workspace.",
            ],
            [
                "No loan moves directly from this screen.",
                "The dashboard is a monitoring layer for the administrative side of the workflow.",
            ],
            [
                "`GET /admin/dashboard` routes non-Admin staff to `AdminService.dashboard(currentSaccoId, principal.getMemberId())`.",
                "`GET /admin/dashboard/database-utilization` returns the JSON payload that feeds the storage-utilization chart.",
            ],
        ),
        screen(
            "4.2",
            "Incidents",
            "docs/assets/lms-workflow-screens/minor-admin/incidents.png",
            [
                "The incidents page combines incident tracking with admin reply and broadcast tools for the active workspace.",
                "It surfaces support messages that members have sent from the Support screen and lets the admin act on them inside the same SACCO scope.",
            ],
            [
                "Updating an incident writes severity, status, and optional resolution notes.",
                "Replying sends a message to a selected member, while broadcasting pushes a message to active members in the current SACCO.",
            ],
            [
                "This is not part of the loan approval ladder, but it affects support and operational follow-up around loan issues.",
                "The next step after an incident change is usually a member notification rather than a workflow status transition on the loan itself.",
            ],
            [
                "`GET /admin/incidents` loads incident rows through `AdminService.incidents(...)`.",
                "`POST /admin/incidents/{id}` updates severity and incident status.",
                "`POST /admin/messages/reply` and `POST /admin/messages/broadcast` create member-facing notifications inside the active SACCO.",
            ],
        ),
        screen(
            "4.3",
            "Users And Roles",
            "docs/assets/lms-workflow-screens/minor-admin/users.png",
            [
                "The Users and Roles table lists workspace accounts, staff role assignments, membership label, and current account status.",
                "The query and pagination controls help the admin manage a growing workspace user list without leaving the scoped SACCO.",
            ],
            [
                "Creating a user validates role assignment rules and persists a new local staff account in the current SACCO.",
                "Updating a row changes staff roles or the account status through the edit modal.",
            ],
            [
                "This screen changes access control, not loan state directly.",
                "The practical downstream effect is who can later review manager queues, board queues, or workspace settings.",
            ],
            [
                "`GET /admin/users` renders a paged `AdminService.usersPage(...)` result.",
                "`POST /admin/users` creates a user in the current SACCO.",
                "`POST /admin/users/{id}` updates role and status assignments for an existing account.",
            ],
        ),
        screen(
            "4.4",
            "User Edit Modal",
            "docs/assets/lms-workflow-screens/minor-admin/users-edit-modal.png",
            [
                "The edit modal exposes the role checkboxes and account status controls for one selected workspace user.",
                "It is the focused write surface for role correction, staff reassignment, and status management.",
            ],
            [
                "Submitting the modal calls the update-user flow, validates allowed roles for the acting admin, and persists the chosen status.",
                "The modal updates the member record rather than creating a new user.",
            ],
            [
                "No credit approval state changes happen here.",
                "The effect is administrative entitlement change, which controls who can see or act on later workflow queues.",
            ],
            [
                "`POST /admin/users/{id}` is the write endpoint behind this modal.",
                "The update is enforced by `AdminService.updateUser(...)`, which applies role and member-status changes inside the scoped SACCO.",
            ],
        ),
        screen(
            "4.5",
            "Station Registry",
            "docs/assets/lms-workflow-screens/minor-admin/station-registry.png",
            [
                "The station registry shows the station set that belongs to the currently scoped SACCO.",
                "For Minor Admins, this is a narrowed operational view rather than a cross-platform registry.",
            ],
            [
                "Edits to the station list are saved through the SACCO update flow, but Minor Admins can only manage stations for their own SACCO.",
                "This keeps workspace identifiers aligned with member registration and external account lookups.",
            ],
            [
                "No loan approval step happens here.",
                "The value of this screen is keeping station identifiers valid so registration and external financial checks continue to resolve correctly.",
            ],
            [
                "`POST /admin/saccos/{saccoId}` calls `SaccoRegistryService.updateStationsOnly(...)` when the actor is a Minor Admin.",
                "Registration and external eligibility later rely on these station IDs being valid for the SACCO.",
            ],
        ),
        screen(
            "4.6",
            "Loan Application Settings",
            "docs/assets/lms-workflow-screens/minor-admin/settings-loan.png",
            [
                "This screen lists the loan products active for the current SACCO together with guarantor count, savings ratio, insurance rate, interest rate, and repayment limits.",
                "It is the workspace control point for product-aware loan behaviour.",
            ],
            [
                "The screen itself is read-focused until the admin chooses edit or create actions.",
                "From here the admin can change the rules that `LoanWorkflowService.saveDraft(...)` and the eligibility preview later enforce.",
            ],
            [
                "No existing application status is changed immediately by opening the list.",
                "The important effect is on future drafts, financial previews, guarantor requirements, and validation boundaries.",
            ],
            [
                "`GET /admin/settings-controls?section=loan` loads active loan products and workspace settings.",
                "The underlying reads come from `AdminService.loanProducts(...)`, `settings(...)`, and the customized-product existence check.",
            ],
        ),
        screen(
            "4.7",
            "Edit Product Modal",
            "docs/assets/lms-workflow-screens/minor-admin/settings-loan-edit-modal.png",
            [
                "The product modal exposes the exact editable rules that drive product behaviour: guarantors required, savings percentage cap, insurance, interest, repayment months, and activation state.",
                "This is the fine-grained point where workspace policy becomes executable system configuration.",
            ],
            [
                "Saving the modal converts percentage inputs to ratios, validates the product identity, and persists the updated settings.",
                "A workspace can also add a customized product through the related create flow.",
            ],
            [
                "Existing applications are not retroactively re-approved by this save.",
                "The real impact is on subsequent drafts and previews that read current SACCO product configuration.",
            ],
            [
                "`POST /admin/settings-controls/{id}` updates an existing product through `AdminService.updateLoanProduct(...)`.",
                "`POST /admin/settings-controls/customized-product` creates a new SACCO-specific product definition.",
            ],
        ),
        screen(
            "4.8",
            "Board Settings",
            "docs/assets/lms-workflow-screens/minor-admin/settings-board.png",
            [
                "The board settings view controls how many board approvals are required before a loan can leave board review.",
                "This quorum value is a business-critical setting because it determines when `AWAITING_BOARD` can transition to `BOARD_APPROVED` or `BOARD_REJECTED`.",
            ],
            [
                "Submitting the form updates the board-review requirement for the current SACCO.",
                "The value is saved at settings level and later read every time board decisions are evaluated.",
            ],
            [
                "This screen does not alter the state of a specific loan immediately.",
                "It changes the future evaluation threshold that `BoardService.evaluateOutcome(...)` applies.",
            ],
            [
                "`POST /admin/settings-controls/review-rules` calls `AdminService.updateBoardReviewRequirement(...)`.",
                "`BoardService` later reads `SaccoSettings.boardQuorum` to decide when board outcome is final.",
            ],
        ),
        screen(
            "4.9",
            "Workspace Outbox Monitor",
            "docs/assets/lms-workflow-screens/minor-admin/outbox.png",
            [
                "The workspace outbox monitor is the scoped version of the platform outbox screen.",
                "It is used to inspect workflow events generated inside the current SACCO, including guarantor, manager, board, and final-status notifications.",
            ],
            [
                "The main write action is retrying an event that should be republished.",
                "Filters help isolate one loan or one date window before the admin retries the event.",
            ],
            [
                "The outbox mirrors workflow transitions; it does not create them.",
                "Operationally, it answers whether a committed business event was queued and whether a retry is needed.",
            ],
            [
                "`GET /admin/outbox` loads the filtered page.",
                "`POST /admin/outbox/{id}/retry` sets the event back to `NEW` for another publish attempt.",
            ],
        ),
        screen(
            "4.10",
            "Workspace Event Log",
            "docs/assets/lms-workflow-screens/minor-admin/events.png",
            [
                "This audit view narrows operational history to events that matter to administrators monitoring one workspace.",
                "It is the forensic screen for tracing who approved, rejected, retried, invited, or updated something.",
            ],
            [
                "The filter bar is read-only and affects only the query window.",
                "No persistent workflow mutation occurs on this screen.",
            ],
            [
                "The approval chain has already happened by the time an entry appears here.",
                "The screen is useful during support and compliance review because it records the observable history of those actions.",
            ],
            [
                "`GET /admin/events` queries `AdminService.eventEntries(...)` with optional date and actor filters.",
                "The data is sourced from persisted `AuditLog` rows.",
            ],
        ),
        screen(
            "5.1",
            "Member Dashboard",
            "docs/assets/lms-workflow-screens/member/dashboard.png",
            [
                "The member dashboard shows current applications, pending guarantees, active repayment exposure, and application status distribution for the signed-in member.",
                "It is the summary landing page after member authentication succeeds.",
            ],
            [
                "The dashboard itself does not submit a loan.",
                "It routes the member into product selection, application detail, archives, guarantees, reports, and support.",
            ],
            [
                "No workflow state changes occur here.",
                "The dashboard reflects the member's current stage across draft, guarantor, manager, board, and disbursed-loan views.",
            ],
            [
                "`GET /app/dashboard` assembles current applications, archived applications, pending guarantor requests, chart rows, and repayment timeline context.",
                "Repayment summary blocks on this page are parsed from stored loan payment summary JSON when it exists.",
            ],
        ),
        screen(
            "5.2",
            "Loan Products",
            "docs/assets/lms-workflow-screens/member/loan-products.png",
            [
                "The Loan Products page lists the active products for the member's SACCO and exposes the apply or calculator actions for each product.",
                "The product set is scoped by `principal.getSaccoId()` and not chosen freely by the member.",
            ],
            [
                "Selecting Apply sends the member to the new application form for the chosen `loanType`.",
                "If the member already has a review-locked application, the page shows the lock and prevents starting another one.",
            ],
            [
                "No application status is created by simply opening the product list.",
                "The next workflow step is the product-aware draft form, where the first `DRAFT` row can be written.",
            ],
            [
                "`GET /app/loan-products` loads `LoanWorkflowService.listProducts(principal.getSaccoId())`.",
                "The in-progress lock indicator uses `LoanWorkflowService.findApplicationInProgress(...)` to stop overlapping review-stage loans.",
            ],
        ),
        screen(
            "5.3",
            "Loan Calculator Modal",
            "docs/assets/lms-workflow-screens/member/loan-products-calculator-modal.png",
            [
                "The calculator modal gives the member a pre-application estimate of eligibility and repayment values before a draft is saved.",
                "It helps the member choose a realistic amount and tenor using the current SACCO's product rules.",
            ],
            [
                "This modal does not create a loan record by itself.",
                "Its role is to preview values that the form will later persist into the draft's financial snapshot.",
            ],
            [
                "No approval stage begins here.",
                "The member moves next to the actual application form once the estimate looks acceptable.",
            ],
            [
                "The eligibility and example values are based on `EligibilityService` and `FinancialDetailsService`.",
                "When the full form is in use, the same preview pipeline is exposed through `GET /app/loan-applications/external-eligibility-summary` and `POST /app/loan-applications/financial-preview`.",
            ],
        ),
        screen(
            "5.4",
            "New Loan Application With Financial Details",
            "docs/assets/workflow-reference-extracted/page26_1.png",
            [
                "This screenshot shows the product-aware application form after the member has entered amount and tenor and loaded the official SACCO financial details.",
                "The page now displays the calculated fees, interest, principal, repayment amount, guarantor selector, attachment uploader, and applicant declaration in one transaction-ready draft form.",
            ],
            [
                "Saving writes or updates a `DRAFT` loan through `LoanWorkflowService.saveDraft(...)` after validating repayment months, dynamic schema fields, guarantor count, same-SACCO guarantors, and eligibility policy.",
                "At this stage the OTP area is informational because the member must first save the application before the submit path becomes available.",
            ],
            [
                "The normal next step is draft persistence first, then later submission to guarantors or directly to manager review depending on the product's guarantor requirement.",
                "The visible `Choose exactly N guarantor(s)` rule comes from the active product configuration for the member's SACCO.",
            ],
            [
                "`POST /app/loan-applications/financial-preview` returns the SACCO financial snapshot shown in the loan-details table.",
                "`GET /app/loan-applications/external-eligibility-summary` loads savings, ratio, and maximum-allowed values from external account context.",
                "`GET /app/guarantors/search` is the AJAX lookup used when the member searches valid guarantor member numbers.",
            ],
        ),
        screen(
            "5.5",
            "Application Ready For Manager Submission",
            "docs/assets/workflow-reference-extracted/page27_1.png",
            [
                "This detail screen shows the application after all required guarantors have approved, so the status is `ALL GUARANTORS APPROVED` and the final member submission block is now unlocked.",
                "The member can review the approved guarantor rows, the loaded financial details, and the final declaration before handing the application to manager review.",
            ],
            [
                "Submitting from this state validates the applicant OTP, records the applicant signature text and verification timestamp, and calls `LoanWorkflowService.submitToManager(...)`.",
                "If the member clicks cancel while the loan is still at manager entry stage, the system can route the removal request through the reversal flow instead of silently discarding a review-stage item.",
            ],
            [
                "The visible next status is `READY_FOR_MANAGER`, which the UI labels as `ON REVIEW BY MANAGER`.",
                "Outbox events are queued so the manager-side workflow can react to the handoff.",
            ],
            [
                "`POST /app/loan-applications/request-signature-otp` emails the applicant confirmation code.",
                "`POST /app/loan-applications/verify-signature-otp` supports client-side OTP validation before the final submit.",
                "`POST /app/loan-applications/{id}/submit` is the write endpoint that advances the loan to manager review.",
            ],
        ),
        screen(
            "5.6",
            "My Applications",
            "docs/assets/lms-workflow-screens/member/loan-applications.png",
            [
                "The My Applications list is the member's registry of current drafts and submitted applications that are not yet archived.",
                "It surfaces the current workflow stage and gives the member a direct path into each application's detail page.",
            ],
            [
                "This list page is read-only until the user opens a record.",
                "Actions such as edit, submit, cancel, or delete are handled on the detail or draft form, not on the list itself.",
            ],
            [
                "No status changes occur here.",
                "The value of the screen is quick status inspection and navigation into the exact application that needs attention.",
            ],
            [
                "`GET /app/loan-applications` loads current application rows from `LoanWorkflowService.myApplications(...)` and derives manager-reason hints through `LoanPresentationService.latestManagerReasons(...)`.",
            ],
        ),
        screen(
            "5.7",
            "Loan Application Detail (Draft And OTP State)",
            "docs/assets/workflow-reference-extracted/page31_1.jpg",
            [
                "This detail state shows a member-owned application that is still `DRAFT`, with the declaration block and applicant OTP area visible before the application leaves the member workspace.",
                "Because the financial details are not yet loaded here, the screen explicitly shows that the official deductions section is still empty.",
            ],
            [
                "The member can return to edit the draft, load the financial details, request an OTP, and then submit when the product rules allow it.",
                "If the application remains a draft, `LoanWorkflowService.saveDraft(...)` can continue updating it; if the member submits, `LoanWorkflowService.submit(...)` decides whether the next state is guarantor review or manager review.",
            ],
            [
                "Draft is the last fully editable state.",
                "Once the loan moves to `AWAITING_GUARANTORS` or `READY_FOR_MANAGER`, the workflow becomes progressively more controlled and visible to other roles.",
            ],
            [
                "`GET /app/loan-applications/{id}` loads the applicant row, attachments, guarantors, repayment context, and saved signature state.",
                "`POST /app/loan-applications/{id}/submit` is the transition endpoint behind the visible submit block.",
            ],
        ),
        screen(
            "5.8",
            "Member Loan Reports",
            "docs/assets/lms-workflow-screens/member/reports.png",
            [
                "The report screen summarizes the member's disbursed, paid, and ongoing loans in a flat reporting table.",
                "It is a historical and analytical view rather than a live approval workspace.",
            ],
            [
                "This screen is read-only.",
                "The main action is exporting the report when the member needs a printable or shareable summary of loan history.",
            ],
            [
                "No workflow transition occurs here.",
                "The report depends on already-recorded disbursement, paid, and defaulted states.",
            ],
            [
                "`GET /app/reports` loads `LoanReportService.memberReport(principal.getMemberId())`.",
                "`GET /documents/reports/member-loans.pdf` generates the downloadable member-loan PDF version of the same reporting dataset.",
            ],
        ),
        screen(
            "5.9",
            "Disbursed And Loan Archive",
            "docs/assets/lms-workflow-screens/member/archives-loans.png",
            [
                "The loan archive shows historical applications that have already left the member's active work queue.",
                "It separates archived loans from current applications so the main list remains focused on actionable records.",
            ],
            [
                "This archive view is read-only and filter-based.",
                "The member uses it to inspect past disbursed, rejected, or closed application records without interfering with active items.",
            ],
            [
                "Archived rows are the output of earlier workflow completion, rejection, or settlement.",
                "No approval chain runs from this screen; it is the evidence side of the workflow.",
            ],
            [
                "`GET /app/archives?section=loans` composes archived rows from `LoanWorkflowService.myApplications(...)` and archive filters.",
            ],
        ),
        screen(
            "5.10",
            "Guarantee Archive",
            "docs/assets/lms-workflow-screens/member/archives-guarantees.png",
            [
                "The guarantee archive stores past guarantor decisions separately from the active guarantee-request queue.",
                "It lets the member review which guarantee commitments were approved, rejected, or otherwise completed historically.",
            ],
            [
                "The archive is read-only.",
                "Removal requests and active approval actions are handled on the live Guarantee Requests screen instead.",
            ],
            [
                "No new guarantor approval decision is made here.",
                "The archive documents completed guarantor history after the live action window has passed.",
            ],
            [
                "`GET /app/archives?section=guarantees` builds the guarantor archive from `LoanWorkflowService.myGuarantorRequests(...)` and archive-specific filters.",
            ],
        ),
        screen(
            "5.11",
            "Guarantee Requests",
            "docs/assets/lms-architecure/screenshots/guarantee-request.png",
            [
                "This is the guarantor's active work queue, showing the applicant, loan type, amount, current status, and actions for a pending guarantee request.",
                "The visible Approve and Reject buttons are the entry points into the guarantor decision flow.",
            ],
            [
                "Approval requires the guarantor to confirm the declaration, request an OTP, validate it, and then call `LoanWorkflowService.approveGuarantorRequest(...)` with signature text and verified time.",
                "Rejection immediately records the refusal, which blocks the application from reaching `ALL_GUARANTORS_APPROVED` until the applicant replaces or resolves the guarantor requirement.",
            ],
            [
                "When all required guarantor requests reach `APPROVED`, the application moves from `AWAITING_GUARANTORS` to `ALL_GUARANTORS_APPROVED`.",
                "That transition unlocks the applicant's final submit-to-manager step.",
            ],
            [
                "`GET /app/guarantee-requests` renders the queue.",
                "`POST /app/guarantee-requests/request-signature-otp` issues the guarantor OTP.",
                "`POST /app/guarantee-requests/{requestId}/approve` and `POST /app/guarantee-requests/{requestId}/reject` record the guarantor decision.",
            ],
        ),
        screen(
            "5.12",
            "Support",
            "docs/assets/lms-workflow-screens/member/support.png",
            [
                "The support form lets a member report technical, account, or workflow issues directly to workspace administration.",
                "It is the member-facing intake point that feeds the admin incidents and message-handling workflow.",
            ],
            [
                "Submitting the form creates a support message for the current SACCO admin team.",
                "The member supplies subject and message text only; the SACCO context is inherited from the authenticated principal.",
            ],
            [
                "This does not alter loan approval state directly.",
                "The next operational step is usually an incident review or an admin reply from the administration workspace.",
            ],
            [
                "`GET /app/support` renders the form.",
                "`POST /app/support` calls `AdminService.submitSupport(principal.getSaccoId(), principal.getMemberId(), subject, message)` to persist and route the support item.",
            ],
        ),
        screen(
            "6.1",
            "Manager Dashboard",
            "docs/assets/lms-workflow-screens/manager/dashboard.png",
            [
                "The manager dashboard summarizes how many loans are disbursed, awaiting repayment completion, waiting in manager review, or already returned and paid.",
                "It also visualizes status distribution so the manager can see queue pressure and recent disbursement activity quickly.",
            ],
            [
                "The dashboard itself is read-only.",
                "The manager uses it to jump into the queue or reports rather than making a decision directly from the chart.",
            ],
            [
                "No application status changes on this page.",
                "It is the monitoring layer for manager workload and disbursement trend visibility.",
            ],
            [
                "`GET /manager/dashboard` loads `ManagerService.dashboard(principal.getSaccoId())`.",
                "The summary is derived from SACCO-scoped `LoanApplication` rows and recent disbursement dates.",
            ],
        ),
        screen(
            "6.2",
            "Manager Queue (On Review By Manager)",
            "docs/assets/lms-workflow-screens/manager/queue-ready.png",
            [
                "The manager queue filtered to `ON REVIEW BY MANAGER` lists applications that have reached `READY_FOR_MANAGER` and are waiting for an accept or reject decision.",
                "Each row represents a member application that has already satisfied the guarantor stage when guarantors were required.",
            ],
            [
                "Opening a row sends the manager into the full review detail page.",
                "No decision is persisted until the manager posts accept or reject from the detail screen.",
            ],
            [
                "The next possible states are `MANAGER_REJECTED` or `AWAITING_BOARD` after the accept path creates board assignments.",
                "This queue is therefore the last purely manager-owned decision gate before board review begins.",
            ],
            [
                "`GET /manager/loan-applications?status=READY_FOR_MANAGER` loads `ManagerService.queue(saccoId, READY_FOR_MANAGER)`.",
            ],
        ),
        screen(
            "6.3",
            "Manager Queue (On Review By Board)",
            "docs/assets/lms-workflow-screens/manager/queue-board.png",
            [
                "This queue view switches the manager's list to applications that already moved beyond manager acceptance and are now in board review.",
                "It helps the manager track items that have left the direct manager queue but still need finalization later.",
            ],
            [
                "The list itself does not write data.",
                "The manager opens a row when they need to inspect the current board-stage detail or prepare for later finalization.",
            ],
            [
                "The next visible progression is from `AWAITING_BOARD` to either `BOARD_APPROVED` or `BOARD_REJECTED`, depending on quorum outcome.",
                "The manager cannot disburse until board approval is reached.",
            ],
            [
                "`GET /manager/loan-applications?status=AWAITING_BOARD` loads the board-stage queue through the same `ManagerService.queue(...)` method with a different status filter.",
            ],
        ),
        screen(
            "6.4",
            "Manager Review Detail",
            "docs/assets/lms-architecure/screenshots/manager-review.png",
            [
                "The manager review detail page consolidates applicant identity, active-loan exposure, financial details, attachments, guarantor rows, and the decision box needed for manager review.",
                "It is the first deep review screen where the manager can confirm the member's data before recording an approval or rejection.",
            ],
            [
                "Rejecting writes a `ManagerReview` row, sets the loan to `MANAGER_REJECTED`, and queues a manager-rejected outbox event for the applicant.",
                "Accepting also writes a manager review row, temporarily sets `MANAGER_ACCEPTED`, assigns the required number of board reviewers, queues board-assignment outbox events, and finally moves the loan to `AWAITING_BOARD`.",
            ],
            [
                "The visible next stage after acceptance is `ON REVIEW BY BOARD`.",
                "If the manager rejects, the workflow returns the loan to the applicant as a rejected item rather than entering board review.",
            ],
            [
                "`GET /manager/loan-applications/{id}` loads the full review model.",
                "`POST /manager/loan-applications/{id}/decision` records ACCEPT or REJECT through `ManagerService.decide(...)`.",
                "`GET /manager/loan-applications/{loanId}/guarantors/{guarantorId}/financial-status` fetches the guarantor savings and shares panel shown when a manager inspects guarantor strength.",
            ],
        ),
        screen(
            "6.5",
            "Manager Review Ready For Disbursement",
            "docs/assets/workflow-reference-extracted/page42_1.png",
            [
                "This reviewed-state screenshot shows the same manager detail page after board quorum has been reached and the loan is ready for final disbursement data entry.",
                "The manager now sees the disbursement block for dates, loan ID, disbursement reference, and repayment notes, alongside the applicant's other active disbursed loans and total exposure.",
            ],
            [
                "Submitting final approval validates the disbursement date, first repayment date, unique numeric loan ID, and schedule inputs, then builds the repayment schedule and marks the loan `FINAL_APPROVED`.",
                "If the board outcome were rejection instead, the final manager action would move the loan to `FINAL_REJECTED` instead of disbursing it.",
            ],
            [
                "The next visible business state after the successful finalize action is the UI label `DISBURSED LOAN`.",
                "That disbursement state is what later enables repayment tracking and payment synchronization.",
            ],
            [
                "`POST /manager/loan-applications/{id}/finalize` calls `ManagerService.finalizeDecision(...)`.",
                "The service builds the repayment schedule through `RepaymentScheduleService`, saves the repayment metadata onto the loan row, and queues a final-status outbox event.",
            ],
        ),
        screen(
            "6.6",
            "Manager Review For A Disbursed Loan",
            "docs/assets/workflow-reference-extracted/page45_1.png",
            [
                "This state shows the manager detail page after disbursement, with the repayment schedule and countdown now visible beneath the core loan summary.",
                "The manager can inspect installments, due dates, payment dates, and current repayment status from the same workflow record.",
            ],
            [
                "At this point the loan is no longer being approved; the manager is monitoring repayment execution and reconciliation.",
                "When needed, the manager can trigger payment synchronization so newly fetched transactions and outstanding-balance summaries are reflected on the record.",
            ],
            [
                "The loan can later move from `FINAL_APPROVED` to `PAID` when the outstanding balance reaches zero after the relevant period, or to `DEFAULTED` when the due date passes with outstanding balance remaining.",
                "Those repayment-state transitions are post-disbursement workflow transitions rather than credit-approval transitions.",
            ],
            [
                "`POST /manager/loan-applications/{id}/sync-payments` starts the payment refresh.",
                "The sync service calls the external member-portal endpoints `/loan-payment-transactions` and `/loan-payment-summary` using member number, station ID, and loan ID.",
                "Fetched rows are saved in `loan_payment_transactions`, while the summary can mark the loan `PAID` or `DEFAULTED` when the repayment conditions are met.",
            ],
        ),
        screen(
            "6.7",
            "Manager Loan Reports",
            "docs/assets/lms-workflow-screens/manager/reports.png",
            [
                "The manager report consolidates disbursed-loan performance, final due dates, paid dates, and portfolio-level return monitoring for the current SACCO.",
                "It is the reporting complement to the transactional review screens.",
            ],
            [
                "The screen is read-only and filter-based by year or returned-only view.",
                "The primary action is exporting the report when the manager needs a PDF portfolio summary.",
            ],
            [
                "No approval or repayment state changes are written from this page.",
                "The report depends on already-recorded final approval, paid, and default states.",
            ],
            [
                "`GET /manager/reports` loads `LoanReportService.managerReport(principal.getSaccoId(), year, returnedOnly)`.",
                "`GET /documents/reports/manager-loans.pdf` exports the same result set as a manager-ready PDF file.",
            ],
        ),
        screen(
            "7.1",
            "Board Queue",
            "docs/assets/lms-workflow-screens/board/queue.png",
            [
                "The board queue lists the applications assigned to the signed-in board member that still have a `PENDING` board decision.",
                "It is a personal assignment queue, not a general list of every board-stage loan in the SACCO.",
            ],
            [
                "Opening a row moves the board member into the review detail page where they can see the committee context and record their decision.",
                "This queue view itself does not persist a board decision.",
            ],
            [
                "The board-stage application remains in `AWAITING_BOARD` until the configured quorum is met for approvals or rejections.",
                "The queue shrinks as the member submits their own decisions.",
            ],
            [
                "`GET /board/queue` filters `BoardService.assignedAll(principal.getMemberId())` down to reviews whose decision is still `PENDING`.",
            ],
        ),
        screen(
            "7.2",
            "Board Review Detail",
            "docs/assets/workflow-reference-extracted/page50_1.jpg",
            [
                "This reviewed-state screenshot shows a board review record after at least one decision has already been recorded and the page is presenting the assessor table, applicant summary, and current board-stage status.",
                "The screen makes the committee's recorded decisions visible in one place so each assessor can see how the review evolved.",
            ],
            [
                "When the review is still pending for the current assessor, the same route can expose the approve or reject form and the board OTP controls; once the decision is already recorded, the page becomes an audit-style detail view.",
                "The guarantor financial-status buttons allow the board member to pull savings and shares information for each guarantor on demand.",
            ],
            [
                "The application remains in board workflow until `BoardService.evaluateOutcome(...)` sees enough approvals or rejections to satisfy quorum.",
                "Once quorum is reached, the loan moves to `BOARD_APPROVED` or `BOARD_REJECTED`, and the manager later handles final approval or final rejection.",
            ],
            [
                "`GET /board/loan-applications/{id}` loads applicant details, committee assessors, attachments, and guarantor rows for the board assignee.",
                "`GET /board/loan-applications/{loanId}/guarantors/{guarantorId}/financial-status` is the fetch behind the guarantor status buttons.",
                "When the current assessor still has a pending decision, `POST /board/loan-applications/{id}/request-signature-otp` and `POST /board/loan-applications/{id}/decision` drive the OTP-backed approval flow.",
            ],
        ),
        screen(
            "7.3",
            "Board Archive",
            "docs/assets/lms-workflow-screens/board/archive.png",
            [
                "The board archive keeps completed or previously decided board assignments separate from the live board queue.",
                "Each row lets the assessor review how they voted and which applications have already left their active queue.",
            ],
            [
                "This screen is read-only.",
                "Board members use it for history lookup rather than for fresh decisions.",
            ],
            [
                "By the time a row appears here, the current assessor's decision is no longer pending.",
                "The broader loan may still be waiting on other assessors, already reviewed, or already finalized by the manager, depending on quorum outcome and later actions.",
            ],
            [
                "`GET /board/archive` loads the same assigned-review set as the queue, but filters out `PENDING` decisions.",
            ],
        ),
        screen(
            "8.1",
            "Printed Loan Application",
            "docs/assets/workflow-reference-extracted/page48_1.png",
            [
                "The printable loan application is the consolidated record of the approved workflow state, showing the application header, SACCO financial details, repayment summary, repayment timetable, guarantor signatures, board-assessor signatures, and the applicant signature block.",
                "This output is intended for document circulation, audit, or archival after the workflow has progressed far enough to make the printable record meaningful.",
            ],
            [
                "The print action does not alter the workflow state.",
                "The system only allows printing when financial details exist and the required guarantor approvals are already satisfied; draft and early guarantor-stage records are blocked from printing.",
            ],
            [
                "The printed document is an end-product view of the workflow, not a new approval stage.",
                "It reflects whichever approval and repayment metadata were already committed onto the loan row at print time.",
            ],
            [
                "`GET /documents/loan-applications/{loanId}/print` builds the HTML document through `LoanPresentationService.buildPrintableHtml(...)`.",
                "The printable generator loads applicant data, guarantor request rows, board review rows, manager reasons, financial snapshot data, and signature verification timestamps before rendering the output.",
            ],
        ),
    ]

    return intro_sections, major_sections, screens


def refine_sections_for_user_document(
    intro_sections: list[TextSection],
    screens: list[ScreenSection],
) -> None:
    for section in intro_sections:
        section.paragraphs = [simplify_user_copy(text) for text in section.paragraphs]
        section.bullets = [simplify_user_copy(text) for text in section.bullets]

    for section in screens:
        section.what_happens = [simplify_user_copy(text) for text in section.what_happens]
        section.on_submit = [simplify_user_copy(text) for text in section.on_submit]
        section.approval_flow = [simplify_user_copy(text) for text in section.approval_flow]
        section.api_communication = EXTERNAL_COMMUNICATION_BY_SECTION.get(section.number, [])


def set_document_defaults(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(31, 41, 55)

    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(12.5)
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 3"].font.size = Pt(11)
    doc.styles["Heading 3"].font.bold = True

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_toc(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = r'TOC \o "1-3" \h \z \u'
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and update the table of contents in Word if fields do not refresh automatically."
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(placeholder)
    run._r.append(fld_char_end)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SACCO Loan Management System Workflow Reference")
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("User Workflow Edition")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(59, 130, 246)

    meta_lines = [
        "Source document reviewed: restructured_lms_workflow_v2.pdf",
        "Generated on: 2026-04-23",
        "Prepared as an export-ready Word/PDF guide with clear workflow explanations for every screenshot.",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line)

    doc.add_page_break()


def add_center_title(doc: Document, text: str, size: float = 16) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 41, 55)


def add_intro_sections(doc: Document, intro_sections: Iterable[TextSection]) -> None:
    doc.add_heading("1. Document Overview", level=1)
    for section in intro_sections:
        doc.add_heading(f"{section.number} {section.title}", level=2)
        for paragraph in section.paragraphs:
            doc.add_paragraph(paragraph)
        for bullet in section.bullets:
            doc.add_paragraph(bullet, style="List Bullet")


def fit_image(image_path: Path) -> tuple[float, float]:
    max_width = 6.6
    max_height = 8.1
    with Image.open(image_path) as img:
        width, height = img.size
    if width == 0 or height == 0:
        return max_width, min(4.0, max_height)
    ratio = min(max_width / width, max_height / height)
    return width * ratio, height * ratio


def add_image(doc: Document, image_path: Path) -> None:
    width, _ = fit_image(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def add_label(doc: Document, label: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(label)
    run.bold = True


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_screen_section(doc: Document, section: ScreenSection) -> None:
    doc.add_heading(f"{section.number} {section.title}", level=2)
    add_image(doc, section.image)
    add_label(doc, "Workflow Explanation:")
    add_label(doc, "What happens:")
    add_bullets(doc, section.what_happens)
    add_label(doc, "On submit:")
    add_bullets(doc, section.on_submit)
    add_label(doc, "Approval flow:")
    add_bullets(doc, section.approval_flow)
    if section.api_communication:
        add_label(doc, "External system communication:")
        add_bullets(doc, section.api_communication)


def main() -> None:
    intro_sections, major_sections, screens = build_sections()
    refine_sections_for_user_document(intro_sections, screens)

    missing = [section.image for section in screens if not section.image.exists()]
    if missing:
        missing_lines = "\n".join(str(path) for path in missing)
        raise FileNotFoundError(f"Missing screenshot assets:\n{missing_lines}")

    doc = Document()
    set_document_defaults(doc)

    footer = doc.sections[0].footer.paragraphs[0]
    add_page_number(footer)

    add_title_page(doc)
    add_center_title(doc, "Table of Contents")
    add_toc(doc)
    doc.add_page_break()

    add_intro_sections(doc, intro_sections)

    by_major = {}
    for item in screens:
        major = item.number.split(".")[0]
        by_major.setdefault(major, []).append(item)

    for major_number, major_title in major_sections:
        doc.add_heading(f"{major_number}. {major_title}", level=1)
        for item in by_major.get(major_number, []):
            add_screen_section(doc, item)

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
